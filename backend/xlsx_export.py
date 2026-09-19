"""Formatted .xlsx workbook builder for the lender pack."""
from io import BytesIO
from typing import List, Optional

from pydantic import BaseModel, Field
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DARK = "FF111827"
DARK_ALT = "FF1E293B"
OLIVE = "FFEEF3E2"
GREEN = "FF7AC24A"
GREEN_TINT = "FFF2F9EC"
ZEBRA = "FFF8FAFC"

THIN = Side(style="thin", color="FFE2E8F0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '"$"#,##0.00'
MONEY0 = '"$"#,##0'
PCT = "0.0%"


class ExportItem(BaseModel):
    name: str = ""
    category_id: Optional[str] = None
    price: float = 0
    unit_cost: float = 0
    units: List[float] = Field(default_factory=list)


class ExportCategory(BaseModel):
    id: str
    name: str = ""


class ExportRequest(BaseModel):
    business_name: str = ""
    industry: str = ""
    months: int = 12
    labels: List[str] = Field(default_factory=list)
    categories: List[ExportCategory] = Field(default_factory=list)
    items: List[ExportItem] = Field(default_factory=list)


def _header(ws, row: int, first_col_title: str, unit_title: str, labels: List[str], total_title: str):
    headers = ["Sales category", first_col_title, unit_title, *labels, total_title]
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=text)
        cell.font = Font(bold=True, color="FFFFFFFF", size=9)
        cell.fill = PatternFill("solid", fgColor=DARK_ALT if c <= 3 + len(labels) else DARK)
        cell.alignment = Alignment(horizontal="left" if c <= 2 else "right", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 28


def _title(ws, row: int, text: str, note: str, width: int):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(3, width))
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(bold=True, color="FFFFFFFF", size=12)
    cell.fill = PatternFill("solid", fgColor=DARK)
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 26
    for c in range(1, max(3, width) + 1):
        ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor=DARK)
    ws.cell(row=row + 1, column=1, value=note).font = Font(italic=True, size=9, color="FF64748B")


def _sizes(ws, months: int):
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 15
    for i in range(months):
        ws.column_dimensions[get_column_letter(4 + i)].width = 11
    ws.column_dimensions[get_column_letter(4 + months)].width = 18


def _sheet(wb, req: ExportRequest, sales: bool):
    months = req.months
    labels = req.labels or [f"M{i + 1}" for i in range(months)]
    ws = wb.create_sheet("Sales Assumptions" if sales else "Cost Assumptions")
    width = 4 + months
    _sizes(ws, months)

    _title(
        ws,
        1,
        ("MONTHLY SALES PROJECTION" if sales else "MONTHLY DIRECT COST PROJECTION")
        + (f" — {req.business_name}" if req.business_name else ""),
        "Enter quantity of units sold in the months you will be paid."
        if sales
        else "Cost per unit x units sold = direct cost per month.",
        width,
    )
    _header(
        ws,
        3,
        "Name of product or service",
        "Sales price ($)" if sales else "Cost per unit ($)",
        labels,
        f"Annual {'revenue' if sales else 'cost'} ($)",
    )

    row = 4
    month_totals = [0.0] * months
    grand_total = 0.0
    revenue_totals = [0.0] * months

    for cat in req.categories:
        cat_items = [i for i in req.items if i.category_id == cat.id]
        if not cat_items:
            continue
        first_row = row
        for n, item in enumerate(cat_items):
            unit = item.price if sales else item.unit_cost
            zebra = PatternFill("solid", fgColor=ZEBRA) if n % 2 else None
            ws.cell(row=row, column=2, value=item.name).alignment = Alignment(horizontal="left")
            price_cell = ws.cell(row=row, column=3, value=unit)
            price_cell.number_format = MONEY
            line_total = 0.0
            for m in range(months):
                u = item.units[m] if m < len(item.units) else 0
                value = u if sales else u * unit
                cell = ws.cell(row=row, column=4 + m, value=value)
                cell.number_format = "#,##0" if sales else MONEY0
                month_totals[m] += (u * unit) if sales else value
                revenue_totals[m] += u * item.price
                line_total += (u * unit) if sales else value
            total_cell = ws.cell(row=row, column=4 + months, value=line_total)
            total_cell.number_format = MONEY0
            total_cell.font = Font(bold=True, color="FF3F6420")
            total_cell.fill = PatternFill("solid", fgColor=GREEN_TINT)
            grand_total += line_total
            for c in range(1, width + 1):
                cell = ws.cell(row=row, column=c)
                cell.border = BORDER
                if zebra and c not in (4 + months,):
                    cell.fill = zebra
                if c >= 3:
                    cell.alignment = Alignment(horizontal="right")
            row += 1

        ws.merge_cells(start_row=first_row, start_column=1, end_row=row - 1, end_column=1)
        cat_cell = ws.cell(row=first_row, column=1, value=cat.name)
        cat_cell.font = Font(bold=True, size=9, color="FF2D3B16")
        cat_cell.alignment = Alignment(horizontal="left", vertical="top", indent=1)
        for r in range(first_row, row):
            c = ws.cell(row=r, column=1)
            c.fill = PatternFill("solid", fgColor=OLIVE)
            c.border = Border(left=Side(style="thick", color=GREEN), right=THIN, top=THIN, bottom=THIN)

    def total_row(label: str, values, fmt: str, fill: str, font_color: str, total_value):
        nonlocal row
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=9, color=font_color)
        for m, v in enumerate(values):
            cell = ws.cell(row=row, column=4 + m, value=v)
            cell.number_format = fmt
        cell = ws.cell(row=row, column=4 + months, value=total_value)
        cell.number_format = fmt
        cell.font = Font(bold=True, size=10, color=font_color)
        for c in range(1, width + 1):
            cc = ws.cell(row=row, column=c)
            cc.fill = PatternFill("solid", fgColor=fill)
            cc.border = BORDER
            if c >= 3:
                cc.alignment = Alignment(horizontal="right")
            if not cc.font.bold:
                cc.font = Font(bold=True, size=9, color=font_color)
        ws.row_dimensions[row].height = 20
        row += 1

    row += 1
    total_row(
        f"TOTAL {'REVENUE' if sales else 'DIRECT COSTS'} ($)",
        month_totals,
        MONEY0,
        DARK,
        "FFFFFFFF",
        grand_total,
    )

    if not sales:
        gp = [revenue_totals[m] - month_totals[m] for m in range(months)]
        revenue_grand = sum(revenue_totals)
        total_row("GROSS PROFIT ($)", gp, MONEY0, OLIVE, "FF2D3B16", revenue_grand - grand_total)
        margins = [(gp[m] / revenue_totals[m]) if revenue_totals[m] else 0 for m in range(months)]
        total_row(
            "GROSS MARGIN (%)",
            margins,
            PCT,
            OLIVE,
            "FF3F6420",
            (revenue_grand - grand_total) / revenue_grand if revenue_grand else 0,
        )

    ws.freeze_panes = "D4"
    return ws


def _cover(wb, req: ExportRequest):
    ws = wb.create_sheet("Summary", 0)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 42

    ws.merge_cells("A1:B1")
    title = ws["A1"]
    title.value = "SAT — SALES & COST ASSUMPTIONS"
    title.font = Font(bold=True, size=14, color="FFFFFFFF")
    title.fill = PatternFill("solid", fgColor=DARK)
    title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws["B1"].fill = PatternFill("solid", fgColor=DARK)

    revenue = sum(
        (i.units[m] if m < len(i.units) else 0) * i.price for i in req.items for m in range(req.months)
    )
    cost = sum(
        (i.units[m] if m < len(i.units) else 0) * i.unit_cost for i in req.items for m in range(req.months)
    )
    margin = (revenue - cost) / revenue if revenue else 0

    rows = [
        ("Business", req.business_name or "—", None),
        ("Industry", req.industry or "—", None),
        ("Projection horizon", f"{req.months} months", None),
        ("First month", req.labels[0] if req.labels else "—", None),
        ("Products / services", len(req.items), "#,##0"),
        ("Sales categories", len([c for c in req.categories if any(i.category_id == c.id for i in req.items)]), "#,##0"),
        ("Total revenue", revenue, MONEY0),
        ("Total direct costs", cost, MONEY0),
        ("Gross profit", revenue - cost, MONEY0),
        ("Gross margin", margin, PCT),
    ]
    for r, (label, value, fmt) in enumerate(rows, start=3):
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = Font(bold=True, size=9, color="FF475569")
        vc = ws.cell(row=r, column=2, value=value)
        vc.font = Font(bold=True, size=11, color="FF0F172A")
        if fmt:
            vc.number_format = fmt
        if label == "Gross margin":
            vc.font = Font(bold=True, size=13, color="FF3F6420")
            vc.fill = PatternFill("solid", fgColor=GREEN_TINT)
        lc.border = BORDER
        vc.border = BORDER

    ws.cell(
        row=len(rows) + 5,
        column=1,
        value="Figures are owner-supplied assumptions for cash flow modelling. Direct costs exclude overheads.",
    ).font = Font(italic=True, size=9, color="FF94A3B8")
    return ws


def build_workbook(req: ExportRequest) -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    _cover(wb, req)
    _sheet(wb, req, sales=True)
    _sheet(wb, req, sales=False)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
